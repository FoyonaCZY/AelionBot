"""Import UTF-8 CSV literally into a formatted workbook (no formula execution).
Original AelionBot helper, MIT; see AELION-LICENSE.txt.
"""
import argparse
import csv
import json
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

def convert(source, target):
    source, target = Path(source).resolve(strict=True), Path(target).resolve()
    if target.suffix.lower() != '.xlsx':
        raise ValueError('Output must have the .xlsx extension')
    if target.exists():
        raise FileExistsError(f'Refusing to overwrite {target}')
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Data'
    widths = {}
    with source.open(encoding='utf-8-sig', newline='') as stream:
        for row_number, row in enumerate(csv.reader(stream), 1):
            if row_number > 100000:
                raise ValueError('Starter importer supports up to 100000 rows; use a streaming workflow for larger files')
            if len(row) > 16384:
                raise ValueError('Input exceeds Excel column limit')
            for column, value in enumerate(row, 1):
                if len(value) > 32767:
                    raise ValueError(f'Cell {row_number}:{column} exceeds Excel text limit')
                cell = sheet.cell(row_number, column, value)
                cell.data_type = 's'  # Even strings beginning with = are literal source data.
                cell.number_format = '@'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if row_number == 1:
                    cell.fill = PatternFill('solid', fgColor='146B70')
                    cell.font = Font(name='Noto Sans', bold=True, color='FFFFFF')
                widths[column] = min(48, max(widths.get(column, 10), len(value) + 2))
    if not widths:
        raise ValueError('CSV is empty')
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = '1:1'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open('xb') as stream:
        workbook.save(stream)
    check = load_workbook(target, read_only=True)
    result = {'path': str(target), 'rows': check.active.max_row, 'columns': check.active.max_column,
              'cellPolicy': 'literal text; convert numeric/date columns deliberately'}
    check.close()
    return result

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input')
    parser.add_argument('output')
    args = parser.parse_args()
    print(json.dumps(convert(args.input, args.output)))
