"""Create an editable 16:9 starter deck from JSON using python-pptx.
Original AelionBot helper, MIT; see AELION-LICENSE.txt.
"""
import argparse
import json
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

INK = '20323D'
ACCENT = '146B70'

def text(slide, x, y, width, height, lines, size=22, color=INK):
    frame = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height)).text_frame
    frame.word_wrap = True
    for i, line in enumerate(lines):
        if not isinstance(line, str):
            raise ValueError('Text entries must be strings')
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.font.name = 'Noto Sans'
        p.font.size = Pt(size)
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(18)
    return frame

def body_lines(value, max_lines=5, max_length=115):
    if not isinstance(value, list) or len(value) > max_lines or any(not isinstance(v, str) or len(v) > max_length for v in value):
        raise ValueError('Body is too dense for the starter layout; split the slide or design a custom layout')
    return value

def create(spec_path, target):
    spec_path, target = Path(spec_path).resolve(strict=True), Path(target).resolve()
    if target.suffix.lower() != '.pptx':
        raise ValueError('Output must have the .pptx extension')
    if target.exists():
        raise FileExistsError(f'Refusing to overwrite {target}')
    spec = json.loads(spec_path.read_text(encoding='utf-8-sig'))
    slides = spec.get('slides')
    if not isinstance(slides, list) or not 1 <= len(slides) <= 100:
        raise ValueError('Provide 1-100 slides')
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    deck.core_properties.title = spec.get('title', '')
    for number, item in enumerate(slides, 1):
        title = item.get('title')
        if not isinstance(title, str) or not title or len(title) > 75:
            raise ValueError('Each slide needs a concise title (up to 75 characters)')
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string('F7FAFA')
        text(slide, .8, .55, 11.7, 1.35, [title], 34)
        if 'columns' in item:
            columns = item['columns']
            if not isinstance(columns, list) or len(columns) != 2:
                raise ValueError('The columns layout requires exactly two columns')
            for i, column in enumerate(columns):
                heading = column.get('heading', '')
                if not isinstance(heading, str) or len(heading) > 32:
                    raise ValueError('Column heading must fit on one line')
                text(slide, .8 + i * 6.0, 2.15, 5.2, .65, [heading], 25, ACCENT)
                text(slide, .8 + i * 6.0, 3.0, 5.2, 3.6, body_lines(column.get('body', []), 4, 80), 22)
        elif 'image' in item:
            image = (spec_path.parent / item['image']).resolve(strict=True)
            picture = slide.shapes.add_picture(str(image), Inches(6.8), Inches(2.1), width=Inches(5.5))
            if picture.height > Inches(4.5):
                ratio = Inches(4.5) / picture.height
                picture.width = int(picture.width * ratio)
                picture.height = Inches(4.5)
            text(slide, .8, 2.25, 5.3, 4.3, body_lines(item.get('body', []), 4, 80), 22)
        else:
            text(slide, .8, 2.25, 11.6, 4.3, body_lines(item.get('body', [])), 25)
        text(slide, 11.6, 6.95, .8, .3, [str(number)], 11, ACCENT)
        slide.notes_slide.notes_text_frame.text = item.get('notes', '')
    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open('xb') as stream:
        deck.save(stream)
    saved = Presentation(target)
    return {'path': str(target), 'slides': len(saved.slides), 'editable': True,
            'next': 'Open in Impress and inspect the PDF export; bounds checks do not replace visual QA'}

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('spec')
    parser.add_argument('output')
    args = parser.parse_args()
    print(json.dumps(create(args.spec, args.output)))
