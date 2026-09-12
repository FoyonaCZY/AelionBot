"""Portable helper regression tests. Run with a Python containing openpyxl and python-pptx."""
import sys
sys.dont_write_bytecode = True
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
from openpyxl import load_workbook
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1] / 'assets' / 'skills'
def module(skill, filename):
    spec = importlib.util.spec_from_file_location(filename, ROOT / skill / 'scripts' / (filename + '.py'))
    result = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(result)
    return result

class Helpers(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='aelion-skill-test-')
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_csv_preserves_ids_and_formula_like_values(self):
        source = self.root / 'source.csv'
        source.write_text('ID,Text\n00123,=1+1\n00004,@SUM(A1)\n00005,中文\n', encoding='utf8')
        target = self.root / 'book.xlsx'
        convert = module('spreadsheet-workbench', 'csv_to_xlsx').convert
        convert(source, target)
        book = load_workbook(target)
        self.assertEqual(book.active['A2'].value, '00123')
        self.assertEqual(book.active['B2'].value, '=1+1')
        self.assertEqual(book.active['B2'].data_type, 's')
        self.assertEqual(book.active['B4'].value, '中文')
        self.assertEqual(book.active.freeze_panes, 'A2')
        book.close()
        before = target.read_bytes()
        with self.assertRaises(FileExistsError): convert(source, target)
        self.assertEqual(target.read_bytes(), before)

    def test_csv_profile_reports_missing_and_malformed_rows(self):
        source = self.root / 'source.csv'
        source.write_text('Label,Value\na,0.1\nb,0.2\nc,\nd,NaN\ne\nf,1,extra\n', encoding='utf8')
        result = module('data-analysis', 'profile_csv').profile(source)
        value = result['columns'][1]
        self.assertEqual(result['rows'], 6)
        self.assertEqual(result['malformedRows'], 2)
        self.assertEqual(value['missing'], 2)
        self.assertEqual(value['nonnumeric'], 1)
        self.assertEqual(value['numeric'], 3)
        self.assertEqual(value['total'], '1.3')
        self.assertEqual(value['minimum'], '0.1')

    def test_deck_is_editable_and_retains_notes(self):
        create = module('presentation-design', 'create_deck').create
        spec = ROOT / 'presentation-design' / 'assets' / 'deck.json'
        target = self.root / 'slides.pptx'
        create(spec, target)
        deck = Presentation(target)
        self.assertEqual(len(deck.slides), 2)
        self.assertIn('The main takeaway', '\n'.join(s.text for s in deck.slides[0].shapes if s.has_text_frame))
        self.assertIn('sources', deck.slides[0].notes_slide.notes_text_frame.text)
        with self.assertRaises(FileExistsError): create(spec, target)
        dense = self.root / 'dense.json'
        dense.write_text(json.dumps({'slides':[{'title':'Dense','body':['x' * 116]}]}))
        with self.assertRaises(ValueError): create(dense, self.root / 'dense.pptx')
        self.assertFalse((self.root / 'dense.pptx').exists())

    def test_export_refuses_to_overwrite_before_launching_office(self):
        source = self.root / 'report.html';source.write_text('<p>Example</p>')
        target = self.root / 'report.pdf';target.write_bytes(b'original')
        export = module('pdf-workbench', 'office_export').export
        with self.assertRaises(FileExistsError): export(source, self.root, 'pdf', executable='never-run')
        self.assertEqual(target.read_bytes(), b'original')

if __name__ == '__main__':
    unittest.main()
